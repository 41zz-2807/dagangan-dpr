"""Blueprint API: laporan omset (ringkasan, harian, produk terlaris, metode bayar)."""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from flask import Blueprint, Response, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import func

from dagang.extensions import db
from dagang.models import Payment, Product, Sale, SaleItem, User
from dagang.utils.decorators import load_current_user, roles_required
from dagang.utils.errors import ApiError

bp = Blueprint("reports", __name__, url_prefix="/api/reports")


def _parse_date(value: str | None, field: str) -> date | None:
    if not value:
        return None
    try:
        return datetime.strptime(value.strip(), "%Y-%m-%d").date()
    except ValueError:
        raise ApiError(400, "invalid_date", f"Filter {field} harus format YYYY-MM-DD.")


def _resolve_range(user: User) -> tuple[datetime, datetime, date, date]:
    """Tentukan rentang laporan.

    Owner boleh memilih rentang bebas. Kasir hanya boleh melihat omset
    hariannya sendiri (tanggal dipaksa ke hari ini).
    """
    today = datetime.now().date()
    if user.role != User.ROLE_OWNER:
        start_date = end_date = today
    else:
        end_date = _parse_date(request.args.get("to"), "to") or today
        start_date = _parse_date(request.args.get("from"), "from") or end_date
        if start_date > end_date:
            raise ApiError(400, "invalid_range", "Tanggal awal tidak boleh melewati akhir.")
        if (end_date - start_date).days > 366:
            raise ApiError(400, "invalid_range", "Rentang maksimal 366 hari.")

    start = datetime.combine(start_date, time.min)
    end = datetime.combine(end_date + timedelta(days=1), time.min)
    return start, end, start_date, end_date


def _scoped_sales(user: User, start: datetime, end: datetime):
    query = db.session.query(Sale).filter(
        Sale.status == Sale.STATUS_PAID,
        Sale.voided_at.is_(None),
        Sale.created_at >= start,
        Sale.created_at < end,
    )
    if user.role != User.ROLE_OWNER:
        query = query.filter(Sale.cashier_id == user.id)
    return query


@bp.get("/summary")
@roles_required(User.ROLE_OWNER)
def summary():
    """Ringkasan omset untuk rentang aktif."""
    user = load_current_user()
    start, end, start_date, end_date = _resolve_range(user)
    sales = _scoped_sales(user, start, end)

    totals = sales.with_entities(
        func.count(Sale.id).label("transaksi"),
        func.coalesce(func.sum(Sale.total), 0).label("omset"),
        func.coalesce(func.sum(Sale.discount), 0).label("diskon"),
    ).one()

    transaksi = int(totals.transaksi or 0)
    omset = totals.omset or 0
    diskon = totals.diskon or 0
    rata_rata = (Decimal(str(omset)) / transaksi) if transaksi else Decimal("0")

    produk_terjual = (
        db.session.query(func.coalesce(func.sum(SaleItem.qty), 0))
        .join(SaleItem.sale)
        .filter(
            Sale.status == Sale.STATUS_PAID,
            Sale.voided_at.is_(None),
            Sale.created_at >= start,
            Sale.created_at < end,
            *([Sale.cashier_id == user.id] if user.role != User.ROLE_OWNER else []),
        )
        .scalar()
    )

    by_method = (
        db.session.query(
            Payment.method,
            func.count(Payment.id),
            func.coalesce(func.sum(Payment.amount), 0),
        )
        .join(Payment.sale)
        .filter(
            Sale.status == Sale.STATUS_PAID,
            Sale.voided_at.is_(None),
            Sale.created_at >= start,
            Sale.created_at < end,
            *([Sale.cashier_id == user.id] if user.role != User.ROLE_OWNER else []),
        )
        .group_by(Payment.method)
        .all()
    )

    by_cashier = (
        db.session.query(
            User.username,
            User.full_name,
            func.count(Sale.id),
            func.coalesce(func.sum(Sale.total), 0),
        )
        .join(User, Sale.cashier_id == User.id)
        .filter(
            Sale.status == Sale.STATUS_PAID,
            Sale.voided_at.is_(None),
            Sale.created_at >= start,
            Sale.created_at < end,
        )
        .group_by(User.username, User.full_name)
        .order_by(func.sum(Sale.total).desc())
        .all()
    )

    low_stock = (
        db.session.query(func.count(Product.id))
        .filter(Product.is_active.is_(True), Product.stock <= Product.min_stock)
        .scalar()
    )

    return jsonify(
        {
            "range": {"from": start_date.isoformat(), "to": end_date.isoformat()},
            "omset": str(omset),
            "transaksi": transaksi,
            "produk_terjual": int(produk_terjual or 0),
            "diskon": str(diskon),
            "rata_rata": str(rata_rata.quantize(Decimal("0.01"))),
            "low_stock_count": int(low_stock or 0),
            "by_method": [
                {"method": method, "count": int(count), "amount": str(amount)}
                for method, count, amount in by_method
            ],
            "by_cashier": [
                {
                    "username": username,
                    "display_name": full_name or username,
                    "count": int(count),
                    "amount": str(amount),
                }
                for username, full_name, count, amount in by_cashier
            ],
        }
    )


@bp.get("/daily")
@roles_required(User.ROLE_OWNER)
def daily():
    """Omset harian (deret per tanggal) untuk grafik sederhana."""
    user = load_current_user()
    start, end, start_date, end_date = _resolve_range(user)
    sales = _scoped_sales(user, start, end)

    rows = (
        sales.with_entities(
            func.date(Sale.created_at).label("tanggal"),
            func.count(Sale.id).label("transaksi"),
            func.coalesce(func.sum(Sale.total), 0).label("omset"),
        )
        .group_by(func.date(Sale.created_at))
        .order_by(func.date(Sale.created_at))
        .all()
    )

    items = [
        {
            "date": tanggal.isoformat() if hasattr(tanggal, "isoformat") else str(tanggal),
            "count": int(count),
            "amount": str(amount),
        }
        for tanggal, count, amount in rows
    ]
    total = sum((Decimal(str(item["amount"])) for item in items), Decimal("0"))
    return jsonify(
        {
            "from": start_date.isoformat(),
            "to": end_date.isoformat(),
            "total": str(total),
            "items": items,
        }
    )


@bp.get("/top-products")
@roles_required(User.ROLE_OWNER)
def top_products():
    """Produk terlaris pada rentang aktif."""
    user = load_current_user()
    start, end, _, _ = _resolve_range(user)
    try:
        limit = min(50, max(1, int(request.args.get("limit", 10))))
    except ValueError:
        raise ApiError(400, "invalid_pagination", "limit harus angka.")

    rows = (
        db.session.query(
            SaleItem.product_id,
            SaleItem.product_name,
            func.coalesce(func.sum(SaleItem.qty), 0).label("qty"),
            func.coalesce(func.sum(SaleItem.subtotal), 0).label("amount"),
        )
        .join(SaleItem.sale)
        .filter(
            Sale.status == Sale.STATUS_PAID,
            Sale.voided_at.is_(None),
            Sale.created_at >= start,
            Sale.created_at < end,
            *([Sale.cashier_id == user.id] if user.role != User.ROLE_OWNER else []),
        )
        .group_by(SaleItem.product_id, SaleItem.product_name)
        .order_by(func.sum(SaleItem.qty).desc())
        .limit(limit)
        .all()
    )

    return jsonify(
        {
            "items": [
                {
                    "product_id": product_id,
                    "product_name": product_name,
                    "qty": int(qty),
                    "amount": str(amount),
                }
                for product_id, product_name, qty, amount in rows
            ]
        }
    )

@bp.get("/export")
@roles_required(User.ROLE_OWNER)
def export_report():
    """Export laporan omset ke Excel/PDF."""
    user = load_current_user()
    fmt = (request.args.get("format") or "xlsx").lower()
    start, end, start_date, end_date = _resolve_range(user)
    sales = _scoped_sales(user, start, end)

    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Laporan Omset"
        
        # Header
        ws["A1"] = "WARUNG DPR - Laporan Omset"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:F1")
        
        ws["A2"] = f"Periode: {start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')}"
        ws.merge_cells("A2:F2")
        
        ws["A4"] = "No. Invoice"
        ws["B4"] = "Tanggal"
        ws["C4"] = "Kasir"
        ws["D4"] = "Subtotal"
        ws["E4"] = "Diskon"
        ws["F4"] = "Total"
        
        for col in ["A4", "B4", "C4", "D4", "E4", "F4"]:
            ws[col].font = Font(bold=True)
            ws[col].alignment = Alignment(horizontal="center")
        
        row = 5
        for sale in sales.order_by(Sale.created_at).all():
            ws[f"A{row}"] = sale.invoice_no
            ws[f"B{row}"] = sale.created_at.strftime("%d/%m/%Y %H:%M") if sale.created_at else ""
            ws[f"C{row}"] = sale.cashier.display_name if sale.cashier else (sale.cashier.username if sale.cashier else "-")
            ws[f"D{row}"] = float(sale.subtotal)
            ws[f"E{row}"] = float(sale.discount)
            ws[f"F{row}"] = float(sale.total)
            row += 1
        
        ws[f"A{row}"] = "TOTAL"
        ws[f"A{row}"].font = Font(bold=True)
        ws[f"D{row}"] = f"=SUM(D5:D{row-1})"
        ws[f"E{row}"] = f"=SUM(E5:E{row-1})"
        ws[f"F{row}"] = f"=SUM(F5:F{row-1})"
        for col in ["D", "E", "F"]:
            ws[f"{col}{row}"].font = Font(bold=True)
        
        for col in ["D", "E", "F"]:
            ws.column_dimensions[col].width = 15
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 20
        
        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        
        filename = f"laporan-omset-{start_date.strftime('%Y%m%d')}-{end_date.strftime('%Y%m%d')}.xlsx"
        return send_file(
            bio,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename,
        )
    elif fmt == "pdf":
        # Use WeasyPrint for PDF
        from weasyprint import HTML
        
        rows_html = ""
        total = Decimal("0")
        for sale in sales.order_by(Sale.created_at).all():
            total += Decimal(str(sale.total))
            cashier_name = sale.cashier.display_name if sale.cashier else "-"
            rows_html += f"""
            <tr>
                <td>{sale.invoice_no}</td>
                <td>{sale.created_at.strftime('%d/%m/%Y %H:%M') if sale.created_at else '-'}</td>
                <td>{cashier_name}</td>
                <td class="right">{float(sale.total):,.0f}</td>
            </tr>
            """
        
        html_content = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <style>
                @page {{ size: A4; margin: 15mm; }}
                body {{ font-family: Arial, sans-serif; font-size: 10pt; }}
                h1 {{ font-size: 16pt; margin-bottom: 5px; }}
                .period {{ color: #666; margin-bottom: 20px; }}
                table {{ width: 100%; border-collapse: collapse; }}
                th {{ background: #f5f5f5; padding: 8px; text-align: left; border: 1px solid #ddd; }}
                td {{ padding: 6px 8px; border: 1px solid #ddd; }}
                .right {{ text-align: right; }}
                .total {{ font-weight: bold; background: #f9f9f9; }}
            </style>
        </head>
        <body>
            <h1>WARUNG DPR - Laporan Omset</h1>
            <div class="period">Periode: {start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')}</div>
            <table>
                <thead>
                    <tr>
                        <th>No. Invoice</th>
                        <th>Tanggal</th>
                        <th>Kasir</th>
                        <th class="right">Total</th>
                    </tr>
                </thead>
                <tbody>
                    {rows_html}
                    <tr class="total">
                        <td colspan="3">TOTAL</td>
                        <td class="right">{float(total):,.0f}</td>
                    </tr>
                </tbody>
            </table>
        </body>
        </html>
        """
        pdf = HTML(string=html_content).write_pdf()
        bio = BytesIO(pdf)
        bio.seek(0)
        filename = f"laporan-omset-{start_date.strftime('%Y%m%d')}-{end_date.strftime('%Y%m%d')}.pdf"
        return send_file(bio, mimetype="application/pdf", as_attachment=True, download_name=filename)
    
    raise ApiError(400, "invalid_format", "Format harus xlsx atau pdf")
