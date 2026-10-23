"""Blueprint API: Pengeluaran."""

from __future__ import annotations

from datetime import datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from flask import Blueprint, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import func

from dagang.extensions import db
from dagang.models import Pengeluaran, PengeluaranKategori, User
from dagang.utils.decorators import load_current_user, roles_required
from dagang.utils.errors import ApiError

bp = Blueprint("pengeluaran", __name__, url_prefix="/api/pengeluaran")


# ---------- Kategori ----------

@bp.get("/kategori")
@roles_required(User.ROLE_OWNER, User.ROLE_KASIR)
def list_kategori():
    items = db.session.query(PengeluaranKategori).order_by(PengeluaranKategori.nama.asc()).all()
    return jsonify({"items": [k.to_dict() for k in items]})


@bp.post("/kategori")
@roles_required(User.ROLE_OWNER)
def create_kategori():
    data = request.get_json(silent=True) or {}
    nama = (data.get("nama") or "").strip()
    if not nama:
        raise ApiError(422, "nama_required", "Nama kategori wajib diisi.")
    if db.session.query(PengeluaranKategori).filter(PengeluaranKategori.nama == nama).first():
        raise ApiError(409, "dup_kategori", "Kategori sudah ada.")
    kategori = PengeluaranKategori(nama=nama)
    db.session.add(kategori)
    db.session.commit()
    return jsonify({"kategori": kategori.to_dict()}), 201


@bp.delete("/kategori/<int:kategori_id>")
@roles_required(User.ROLE_OWNER)
def delete_kategori(kategori_id: int):
    kategori = db.session.query(PengeluaranKategori).filter_by(id=kategori_id).first()
    if not kategori:
        raise ApiError(404, "not_found", "Kategori tidak ditemukan.")

    used = (
        db.session.query(func.count(Pengeluaran.id))
        .filter(Pengeluaran.kategori_id == kategori_id)
        .scalar()
        or 0
    )
    if used:
        raise ApiError(
            409,
            "kategori_used",
            f"Kategori masih dipakai oleh {used} pengeluaran. "
            "Hapus atau ubah pengeluaran tersebut terlebih dahulu.",
        )

    db.session.delete(kategori)
    db.session.commit()
    return jsonify({"message": "Kategori dihapus."})


# ---------- Pengeluaran ----------

@bp.get("")
@roles_required(User.ROLE_OWNER, User.ROLE_KASIR)
def list_pengeluaran():
    user = load_current_user()
    query = db.session.query(Pengeluaran)

    if user.role != User.ROLE_OWNER:
        query = query.filter(Pengeluaran.created_by_id == user.id)

    kategori_id = request.args.get("kategori_id")
    if kategori_id:
        try:
            query = query.filter(Pengeluaran.kategori_id == int(kategori_id))
        except ValueError:
            raise ApiError(400, "invalid_filter", "kategori_id harus angka.")

    date_from = request.args.get("from")
    date_to = request.args.get("to")
    if date_from:
        try:
            query = query.filter(Pengeluaran.tanggal >= datetime.strptime(date_from, "%Y-%m-%d"))
        except ValueError:
            raise ApiError(400, "invalid_filter", "Format tanggal 'from' harus YYYY-MM-DD.")
    if date_to:
        try:
            query = query.filter(Pengeluaran.tanggal < datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1))
        except ValueError:
            raise ApiError(400, "invalid_filter", "Format tanggal 'to' harus YYYY-MM-DD.")

    try:
        page = max(1, int(request.args.get("page", 1)))
        per_page = min(200, max(1, int(request.args.get("per_page", 50))))
    except ValueError:
        raise ApiError(400, "invalid_pagination", "page/per_page harus angka.")

    total = query.count()
    items = (
        query.order_by(Pengeluaran.tanggal.desc(), Pengeluaran.id.desc())
        .offset((page - 1) * per_page)
        .limit(per_page)
        .all()
    )
    return jsonify({
        "items": [p.to_dict() for p in items],
        "total": total,
        "page": page,
        "per_page": per_page,
        "pages": (total + per_page - 1) // per_page,
    })


@bp.post("")
@roles_required(User.ROLE_OWNER, User.ROLE_KASIR)
def create_pengeluaran():
    user = load_current_user()
    data = request.get_json(silent=True) or {}

    try:
        tanggal = datetime.strptime(data.get("tanggal", ""), "%Y-%m-%d")
    except (ValueError, TypeError):
        raise ApiError(422, "invalid_tanggal", "Tanggal wajib diisi format YYYY-MM-DD.")

    nominal_raw = data.get("nominal")
    try:
        nominal = int(nominal_raw)
    except (TypeError, ValueError):
        raise ApiError(422, "invalid_nominal", "Nominal harus angka.")
    if nominal <= 0:
        raise ApiError(422, "invalid_nominal", "Nominal harus lebih dari nol.")

    kategori_id = data.get("kategori_id")
    if kategori_id:
        try:
            kategori_id = int(kategori_id)
        except (TypeError, ValueError):
            raise ApiError(422, "invalid_kategori", "kategori_id harus angka.")
        if not db.session.query(PengeluaranKategori).filter_by(id=kategori_id).first():
            raise ApiError(422, "invalid_kategori", "Kategori tidak ditemukan.")

    metode_bayar = (data.get("metode_bayar") or "tunai").strip()
    keterangan = (data.get("keterangan") or "").strip() or None
    kategori_nama = None
    if kategori_id:
        k = db.session.query(PengeluaranKategori).filter_by(id=kategori_id).first()
        if k:
            kategori_nama = k.nama

    p = Pengeluaran(
        tanggal=tanggal,
        kategori_id=kategori_id,
        kategori_nama=kategori_nama,
        nominal=nominal,
        keterangan=keterangan,
        metode_bayar=metode_bayar,
        created_by_id=user.id,
    )
    db.session.add(p)
    db.session.commit()
    return jsonify({"pengeluaran": p.to_dict()}), 201


@bp.get("/<int:pengeluaran_id>")
@roles_required(User.ROLE_OWNER, User.ROLE_KASIR)
def get_pengeluaran(pengeluaran_id: int):
    user = load_current_user()
    p = db.session.query(Pengeluaran).filter_by(id=pengeluaran_id).first()
    if not p:
        raise ApiError(404, "not_found", "Pengeluaran tidak ditemukan.")
    if user.role != User.ROLE_OWNER and p.created_by_id != user.id:
        raise ApiError(403, "forbidden", "Tidak diizinkan.")
    return jsonify({"pengeluaran": p.to_dict()})


@bp.delete("/<int:pengeluaran_id>")
@roles_required(User.ROLE_OWNER, User.ROLE_KASIR)
def delete_pengeluaran(pengeluaran_id: int):
    user = load_current_user()
    p = db.session.query(Pengeluaran).filter_by(id=pengeluaran_id).first()
    if not p:
        raise ApiError(404, "not_found", "Pengeluaran tidak ditemukan.")
    if user.role != User.ROLE_OWNER and p.created_by_id != user.id:
        raise ApiError(403, "forbidden", "Tidak diizinkan.")
    db.session.delete(p)
    db.session.commit()
    return jsonify({"message": "Pengeluaran dihapus."})


@bp.get("/summary")
@roles_required(User.ROLE_OWNER, User.ROLE_KASIR)
def summary():
    user = load_current_user()
    query = db.session.query(Pengeluaran)
    if user.role != User.ROLE_OWNER:
        query = query.filter(Pengeluaran.created_by_id == user.id)

    date_from = request.args.get("from")
    date_to = request.args.get("to")
    if date_from:
        try:
            query = query.filter(Pengeluaran.tanggal >= datetime.strptime(date_from, "%Y-%m-%d"))
        except ValueError:
            raise ApiError(400, "invalid_filter", "Format tanggal 'from' harus YYYY-MM-DD.")
    if date_to:
        try:
            query = query.filter(Pengeluaran.tanggal < datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1))
        except ValueError:
            raise ApiError(400, "invalid_filter", "Format tanggal 'to' harus YYYY-MM-DD.")

    total = query.with_entities(func.coalesce(func.sum(Pengeluaran.nominal), 0)).scalar() or 0
    count = query.count()
    return jsonify({
        "total": int(total),
        "count": count,
    })


@bp.get("/export")
@roles_required(User.ROLE_OWNER, User.ROLE_KASIR)
def export_pengeluaran():
    user = load_current_user()
    fmt = (request.args.get("format") or "xlsx").lower()

    query = db.session.query(Pengeluaran)
    if user.role != User.ROLE_OWNER:
        query = query.filter(Pengeluaran.created_by_id == user.id)

    date_from = request.args.get("from")
    date_to = request.args.get("to")
    if date_from:
        try:
            query = query.filter(Pengeluaran.tanggal >= datetime.strptime(date_from, "%Y-%m-%d"))
        except ValueError:
            raise ApiError(400, "invalid_filter", "Format tanggal 'from' harus YYYY-MM-DD.")
    if date_to:
        try:
            query = query.filter(Pengeluaran.tanggal < datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1))
        except ValueError:
            raise ApiError(400, "invalid_filter", "Format tanggal 'to' harus YYYY-MM-DD.")

    items = query.order_by(Pengeluaran.tanggal.desc(), Pengeluaran.id.desc()).all()

    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Pengeluaran"
        ws["A1"] = "WARUNG DPR - Laporan Pengeluaran"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:F1")
        ws["A2"] = "Daftar pengeluaran"
        ws.merge_cells("A2:F2")

        headers = ["Tanggal", "Kategori", "Nominal", "Keterangan", "Metode Bayar", "Dibuat Oleh"]
        for idx, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=idx, value=h)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        row = 5
        for p in items:
            ws[f"A{row}"] = p.tanggal.strftime("%d/%m/%Y") if p.tanggal else ""
            ws[f"B{row}"] = p.kategori_nama or "-"
            ws[f"C{row}"] = p.nominal
            ws[f"D{row}"] = p.keterangan or "-"
            ws[f"E{row}"] = p.metode_bayar
            ws[f"F{row}"] = p.created_by.display_name if p.created_by else "-"
            row += 1

        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 30
        ws.column_dimensions["E"].width = 15
        ws.column_dimensions["F"].width = 20

        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        fname = f"pengeluaran-{datetime.now().strftime('%Y%m%d')}.xlsx"
        return send_file(bio, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                         as_attachment=True, download_name=fname)

    elif fmt == "pdf":
        from weasyprint import HTML
        rows_html = ""
        total = 0
        for p in items:
            total += p.nominal
            rows_html += f"""<tr>
                <td>{p.tanggal.strftime('%d/%m/%Y') if p.tanggal else '-'}</td>
                <td>{p.kategori_nama or '-'}</td>
                <td class="right">{p.nominal:,}</td>
                <td>{p.keterangan or '-'}</td>
                <td>{p.metode_bayar}</td>
            </tr>"""
        html_content = f"""<!DOCTYPE html><html><head><meta charset="utf-8">
        <style>@page{{size:A4;margin:15mm}}body{{font-family:Arial,sans-serif;font-size:9pt}}
        h1{{font-size:14pt}}table{{width:100%;border-collapse:collapse}}th{{background:#f5f5f5;padding:6px;border:1px solid #ddd;text-align:left}}
        td{{padding:4px 6px;border:1px solid #ddd}}.right{{text-align:right}}.total{{font-weight:bold;background:#f9f9f9}}</style>
        </head><body>
        <h1>WARUNG DPR - Laporan Pengeluaran</h1>
        <table><thead><tr><th>Tanggal</th><th>Kategori</th><th class="right">Nominal</th><th>Keterangan</th><th>Metode Bayar</th></tr></thead>
        <tbody>{rows_html}
        <tr class="total"><td colspan="2">TOTAL</td><td class="right">{total:,}</td><td colspan="2"></td></tr>
        </tbody></table></body></html>"""
        pdf = HTML(string=html_content).write_pdf()
        bio = BytesIO(pdf)
        bio.seek(0)
        fname = f"pengeluaran-{datetime.now().strftime('%Y%m%d')}.pdf"
        return send_file(bio, mimetype="application/pdf", as_attachment=True, download_name=fname)

    raise ApiError(400, "invalid_format", "Format harus xlsx atau pdf")
