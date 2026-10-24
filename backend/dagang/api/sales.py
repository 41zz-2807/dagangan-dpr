"""Blueprint API: transaksi penjualan (buat, daftar, detail)."""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
from io import BytesIO

from flask import Blueprint, Response, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import func
from sqlalchemy import or_

from dagang.extensions import db
from dagang.models import Payment, Sale, SaleItem, User
from dagang.schemas import SaleCreateRequest, SaleVoidRequest
from dagang.services.receipt import build_receipt_pdf
from dagang.services.sale import create_sale
from dagang.services.telegram import send_message
from dagang.utils.decorators import load_current_user, roles_required
from dagang.utils.errors import ApiError
from dagang.utils.validators import parse_body

bp = Blueprint("sales", __name__, url_prefix="/api/sales")


def _parse_date(value: str | None, field: str) -> date | None:
    if not value:
        return None
    try:
        return datetime.strptime(value.strip(), "%Y-%m-%d").date()
    except ValueError:
        raise ApiError(400, "invalid_date", f"Filter {field} harus format YYYY-MM-DD.")


def _day_bounds(start: date | None, end: date | None):
    """Rentang waktu lokal (Asia/Jakarta) dari filter tanggal."""
    begin = datetime.combine(start, time.min) if start else None
    finish = datetime.combine(end + timedelta(days=1), time.min) if end else None
    return begin, finish


@bp.post("")
@roles_required(User.ROLE_OWNER, User.ROLE_KASIR)
def create_sale_endpoint():
    """Simpan transaksi baru (total dihitung ulang di server)."""
    data = parse_body(SaleCreateRequest)
    user = load_current_user()

    sale = create_sale(
        items=data.items,
        payments=data.payments,
        discount=data.discount,
        customer_name=data.customer_name,
        customer_phone=data.customer_phone,
        note=data.note,
        user=user,
    )
    try:
        send_message(
            f"🛒 <b>Penjualan baru</b>\n"
            f"Invoice: {sale.invoice_no}\n"
            f"Total: {sale.total}\n"
            f"Kasir: {user.username}\n"
            f"Metode: {', '.join(p.method for p in sale.payments)}"
        )
    except Exception:
        pass
    return jsonify({"sale": sale.to_dict(detail=True)}), 201


@bp.get("")
@roles_required(User.ROLE_OWNER, User.ROLE_KASIR)
def list_sales():
    """Daftar transaksi. Kasir hanya melihat transaksinya sendiri."""
    user = load_current_user()
    query = db.session.query(Sale)

    if user.role != User.ROLE_OWNER:
        query = query.filter(Sale.cashier_id == user.id)
        # Default kasir: miliknya sendiri hari ini.
        if not request.args.get("from") and not request.args.get("to"):
            today = datetime.now().date()
            query = query.filter(
                Sale.created_at >= datetime.combine(today, time.min),
                Sale.created_at < datetime.combine(today + timedelta(days=1), time.min),
            )
    else:
        cashier_id = request.args.get("cashier_id")
        if cashier_id:
            try:
                query = query.filter(Sale.cashier_id == int(cashier_id))
            except ValueError:
                raise ApiError(400, "invalid_filter", "cashier_id harus angka.")

    start = _parse_date(request.args.get("from"), "from")
    end = _parse_date(request.args.get("to"), "to")
    begin, finish = _day_bounds(start, end)
    if begin is not None:
        query = query.filter(Sale.created_at >= begin)
    if finish is not None:
        query = query.filter(Sale.created_at < finish)

    status = (request.args.get("status") or "").strip().lower()
    if status:
        if status not in Sale.STATUSES:
            raise ApiError(400, "invalid_filter", "status harus 'paid' atau 'void'.")
        query = query.filter(Sale.status == status)

    method = (request.args.get("method") or "").strip().lower()
    if method:
        if method not in Payment.METHODS:
            raise ApiError(400, "invalid_filter", "method tidak dikenal.")
        query = query.filter(
            Sale.id.in_(
                db.session.query(Payment.sale_id).filter(Payment.method == method)
            )
        )

    search = (request.args.get("q") or "").strip()
    if search:
        pattern = f"%{search.lower()}%"
        query = query.filter(
            or_(
                db.func.lower(Sale.invoice_no).like(pattern),
                Sale.id.in_(
                    db.session.query(SaleItem.sale_id)
                    .filter(db.func.lower(SaleItem.product_name).like(pattern))
                ),
            )
        )

    try:
        page = max(1, int(request.args.get("page", 1)))
        per_page = min(200, max(1, int(request.args.get("per_page", 50))))
    except ValueError:
        raise ApiError(400, "invalid_pagination", "page/per_page harus angka.")

    total = query.count()
    sales = (
        query.order_by(Sale.created_at.desc(), Sale.id.desc())
        .offset((page - 1) * per_page)
        .limit(per_page)
        .all()
    )
    return jsonify(
        {
            "items": [sale.to_dict() for sale in sales],
            "total": total,
            "page": page,
            "per_page": per_page,
            "pages": (total + per_page - 1) // per_page,
        }
    )


@bp.get("/<int:sale_id>")
@roles_required(User.ROLE_OWNER, User.ROLE_KASIR)
def get_sale(sale_id: int):
    """Detail satu transaksi."""
    user = load_current_user()
    sale = db.session.get(Sale, sale_id)
    if sale is None:
        raise ApiError(404, "not_found", "Transaksi tidak ditemukan.")
    if user.role != User.ROLE_OWNER and sale.cashier_id != user.id:
        raise ApiError(403, "forbidden", "Anda tidak memiliki izin untuk transaksi ini.")
    return jsonify({"sale": sale.to_dict(detail=True)})


@bp.post("/<int:sale_id>/void")
@roles_required(User.ROLE_OWNER)
def void_sale_endpoint(sale_id: int):
    """Batalkan transaksi (owner): stok dikembalikan, data tidak dihapus."""
    data = parse_body(SaleVoidRequest)
    owner = load_current_user()
    sale = void_sale(sale_id=sale_id, reason=data.reason, user=owner)
    return jsonify({"sale": sale.to_dict(detail=True)})


@bp.get("/<int:sale_id>/receipt.pdf")
@roles_required(User.ROLE_OWNER, User.ROLE_KASIR)
def sale_receipt_pdf(sale_id: int):
    """Unduh struk PDF 80mm milik transaksi ini."""
    user = load_current_user()
    sale = db.session.get(Sale, sale_id)
    if sale is None:
        raise ApiError(404, "not_found", "Transaksi tidak ditemukan.")
    if user.role != User.ROLE_OWNER and sale.cashier_id != user.id:
        raise ApiError(403, "forbidden", "Anda tidak memiliki izin untuk transaksi ini.")

    pdf = build_receipt_pdf(sale)
    response = send_file(
        BytesIO(pdf),
        mimetype="application/pdf",
        as_attachment=False,
        download_name=f"struk-{sale.invoice_no}.pdf",
    )
    response.headers["Content-Disposition"] = f'inline; filename="struk-{sale.invoice_no}.pdf"'
    response.headers["Cache-Control"] = "no-store"
    return response

@bp.get("/export")
@roles_required(User.ROLE_OWNER, User.ROLE_KASIR)
def export_sales():
    """Export daftar transaksi ke Excel/PDF."""
    user = load_current_user()
    fmt = (request.args.get("format") or "xlsx").lower()
    
    from_date = request.args.get("from")
    to_date = request.args.get("to")
    today = datetime.now().date()
    
    if user.role != User.ROLE_OWNER:
        start_date = end_date = today
    else:
        try:
            end_date = datetime.strptime(to_date, "%Y-%m-%d").date() if to_date else today
            start_date = datetime.strptime(from_date, "%Y-%m-%d").date() if from_date else end_date
        except Exception:
            raise ApiError(400, "invalid_date", "Format tanggal harus YYYY-MM-DD")
    
    if start_date > end_date:
        raise ApiError(400, "invalid_range", "Tanggal awal tidak boleh melewati akhir")
    
    start = datetime.combine(start_date, time.min)
    end = datetime.combine(end_date + timedelta(days=1), time.min)
    
    query = db.session.query(Sale).filter(
        Sale.created_at >= start,
        Sale.created_at < end,
    )
    if user.role != User.ROLE_OWNER:
        query = query.filter(Sale.cashier_id == user.id)
    
    sales = query.order_by(Sale.created_at.desc()).all()
    
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Transaksi"
        ws["A1"] = "WARUNG DPR - Daftar Transaksi"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:H1")
        ws["A2"] = f"Periode: {start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')}"
        ws.merge_cells("A2:H2")
        headers = ["No. Invoice", "Tanggal", "Kasir", "Status", "Subtotal", "Diskon", "Total", "Metode Bayar"]
        for idx, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=idx, value=h)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        row = 5
        for sale in sales:
            methods = ", ".join([p.method for p in sale.payments])
            ws[f"A{row}"] = sale.invoice_no
            ws[f"B{row}"] = sale.created_at.strftime("%d/%m/%Y %H:%M") if sale.created_at else ""
            ws[f"C{row}"] = sale.cashier.display_name if sale.cashier and sale.cashier.display_name else (sale.cashier.username if sale.cashier else "-")
            ws[f"D{row}"] = sale.status
            ws[f"E{row}"] = float(sale.subtotal)
            ws[f"F{row}"] = float(sale.discount)
            ws[f"G{row}"] = float(sale.total)
            ws[f"H{row}"] = methods
            row += 1
        for col in ["E", "F", "G"]:
            ws.column_dimensions[col].width = 15
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["H"].width = 25
        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        filename = f"transaksi-{start_date.strftime('%Y%m%d')}-{end_date.strftime('%Y%m%d')}.xlsx"
        return send_file(bio, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=filename)
    elif fmt == "pdf":
        from weasyprint import HTML
        rows_html = ""
        for sale in sales:
            methods = ", ".join([p.method for p in sale.payments])
            cashier_name = sale.cashier.display_name if sale.cashier and sale.cashier.display_name else (sale.cashier.username if sale.cashier else "-")
            rows_html += f"""<tr><td>{sale.invoice_no}</td><td>{sale.created_at.strftime('%d/%m/%Y %H:%M') if sale.created_at else '-'}</td><td>{cashier_name}</td><td>{sale.status}</td><td class="right">{float(sale.total):,.0f}</td><td>{methods}</td></tr>"""
        html_content = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><style>@page{{size:A4;margin:10mm}}body{{font-family:Arial;font-size:8pt}}h1{{font-size:14pt}}table{{width:100%;border-collapse:collapse}}th{{background:#f5f5f5;padding:6px;border:1px solid #ddd}}td{{padding:4px 6px;border:1px solid #ddd}}.right{{text-align:right}}</style></head><body><h1>WARUNG DPR - Daftar Transaksi</h1><div>Periode: {start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')}</div><table><thead><tr><th>Invoice</th><th>Tanggal</th><th>Kasir</th><th>Status</th><th class="right">Total</th><th>Metode</th></tr></thead><tbody>{rows_html}</tbody></table></body></html>"""
        pdf = HTML(string=html_content).write_pdf()
        bio = BytesIO(pdf)
        bio.seek(0)
        filename = f"transaksi-{start_date.strftime('%Y%m%d')}-{end_date.strftime('%Y%m%d')}.pdf"
        return send_file(bio, mimetype="application/pdf", as_attachment=True, download_name=filename)
    raise ApiError(400, "invalid_format", "Format harus xlsx atau pdf")
